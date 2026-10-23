"""The cost-model spreadsheet reproduces Table 6.0's decomposed month from its inputs."""
import pathlib
import shutil
import subprocess
import sys

import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "cost-model"))
import build_cost_model as b  # noqa: E402

EXPECTED_K = {"Inference": 118, "Provisioned": 46, "Embeddings": 9, "Vector": 21, "GPU": 14,
              "FineTuning": 7, "Satellites": 11}


def test_workbook_has_one_tab_per_component(tmp_path):
    from openpyxl import load_workbook
    wb = load_workbook(b.build(tmp_path / "m.xlsx"))
    assert wb.sheetnames == ["Summary"] + list(EXPECTED_K)
    assert wb["Summary"]["B11"].value.startswith("=SUM(")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice recalculates the formulas")
def test_formulas_reproduce_table_6_0(tmp_path):
    from openpyxl import load_workbook
    src = b.build(tmp_path / "m.xlsx")
    out = tmp_path / "calc"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180)
    s = load_workbook(out / "m.xlsx", data_only=True)["Summary"]
    got = {s.cell(row=r, column=1).value: s.cell(row=r, column=2).value for r in range(4, 11)}
    for k, v in EXPECTED_K.items():
        assert round(got[k] / 1000) == v, (k, got[k])
    assert round(s["B11"].value / 1000) == 226
    assert s["B15"].value == pytest.approx(0.615, abs=0.005)        # 61% output tokens
